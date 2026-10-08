"""Guards the downloadable history-import template against drifting from the importer."""
import importlib.util
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from src.utils import excel_to_postgres

SCRIPT_PATH = Path(__file__).resolve().parents[2] / "scripts" / "generate_import_template.py"


def _as_date(cell):
    return cell.date() if isinstance(cell, datetime) else cell


@pytest.fixture(scope="module")
def generator():
    spec = importlib.util.spec_from_file_location("generate_import_template", SCRIPT_PATH)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_committed_template_matches_generator(generator):
    wb = openpyxl.load_workbook(generator.OUTPUT_PATH, data_only=True)
    actual = [
        tuple(_as_date(c) for c in row)
        for row in wb[generator.TRANSACTIONS_SHEET].iter_rows(values_only=True)
    ]
    expected = [tuple(generator.HEADERS)] + [tuple(row) for row in generator.ROWS]
    assert actual == expected, "Template is stale: run python3 importer/scripts/generate_import_template.py"


def test_template_parses_cleanly_with_default_mapping(generator):
    # dry_run=True uses placeholder IDs, so no DB access is needed.
    transactions, skipped, skip_details, _ = excel_to_postgres.process_excel_file(
        str(generator.OUTPUT_PATH), dry_run=True
    )
    assert skipped == 0, skip_details
    assert len(transactions) == len(generator.ROWS)
    assert sum(1 for t in transactions if t.transaction.is_ignored) == 2
    assert {t.category_name for t in transactions if t.category_name} >= {"Income:Salary", "Gift"}


def test_rows_missing_required_values_are_skipped_and_blank_rows_ignored(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Account", "Date", "Payee", "Category", "Amount", "Status", "Notes"])
    ws.append(["Checking", datetime(2026, 1, 5), "Coffee", "Food", -3.5, "Cleared", None])
    ws.append(["Checking", datetime(2026, 1, 6), None, "Food", -9, None, None])
    ws.append([None, None, None, None, None, None, None])
    ws.append(["Checking", None, "Bakery", None, None, None, None])
    path = tmp_path / "rows.xlsx"
    wb.save(path)

    transactions, skipped, skip_details, _ = excel_to_postgres.process_excel_file(str(path), dry_run=True)

    assert len(transactions) == 1
    assert skipped == 2
    assert skip_details == [
        {"row": 3, "reason": "missing required value: Payee"},
        {"row": 5, "reason": "missing required value: Date, Amount"},
    ]
