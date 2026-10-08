"""Generates the history-import Excel template served by the app.

Usage (from the repo root):
    python3 importer/scripts/generate_import_template.py

All data below is fictional. The headers must match the default column mapping
in src/utils/excel_to_postgres.py; tests/utils/test_import_template.py enforces it.
"""
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

OUTPUT_PATH = Path(__file__).resolve().parents[2] / "app" / "client" / "public" / "pfm-import-template.xlsx"

TRANSACTIONS_SHEET = "Transactions"
INSTRUCTIONS_SHEET = "Instructions"

HEADERS = ["Account", "Date", "Payee", "Category", "Amount", "Status", "Notes"]

ROWS = [
    ["Main Checking", date(2026, 1, 2), "Acme Corp", "Income:Salary", 4500.00, "cleared", "January salary"],
    ["Main Checking", date(2026, 1, 3), "Sunrise Apartments", "Housing:Rent", -1350.00, "cleared", None],
    ["Main Checking", date(2026, 1, 5), "Green Market", "Food:Groceries", -87.45, "cleared", None],
    ["Visa Credit Card", date(2026, 1, 6), "Corner Bistro", "Food:Restaurants", -42.80, "cleared", "Dinner with friends"],
    ["Visa Credit Card", date(2026, 1, 9), "StreamFlix", "Leisure:Streaming", -15.99, "cleared", "Monthly subscription"],
    ["Main Checking", date(2026, 1, 12), "City Pharmacy", "Health:Pharmacy", -23.10, "cleared", None],
    ["Visa Credit Card", date(2026, 1, 15), "Fuel Stop", "Transport:Fuel", -60.00, "pending", None],
    ["Main Checking", date(2026, 1, 20), "Transfer to Visa Credit Card", None, -500.00, "ignored", "Card payment, not an expense"],
    ["Visa Credit Card", date(2026, 1, 20), "Payment received", None, 500.00, "ignored", "Card payment, not income"],
    ["Main Checking", date(2026, 1, 28), "Book Nook", "Gift", -34.90, "pending", "No group, so it lands in Others"],
]

INSTRUCTIONS = [
    ("How to use this template", None),
    ("Fill in the Transactions sheet (one row per transaction) and upload it under Settings > History.", None),
    ("Keep the header row as is. Only the first sheet is read; this Instructions sheet is ignored.", None),
    (None, None),
    ("Column", "What to put there"),
    ("Account", "Required. Account name. Accounts that don't exist yet are created for you."),
    ("Date", "Required. A real Excel date cell (or an ISO text like 2026-01-31). Avoid dd/mm/yyyy text: it may be read as mm/dd."),
    ("Payee", "Required. Who you paid or who paid you."),
    ("Category", "Optional value, but the column must exist. Format Group:Category (e.g. Food:Groceries). Without a colon the group is Others. Leave empty for uncategorized."),
    ("Amount", "Required. A number, not text. Negative = money out (expenses), positive = money in (income)."),
    ("Status", "Optional. pending, cleared or reconciled. Use ignored to keep a row in balances but out of reports (e.g. transfers between your own accounts). Anything else is treated as pending."),
    ("Notes", "Optional. Free text."),
    (None, None),
    ("Rows missing Account, Date, Payee or Amount are skipped. Re-importing the same file is safe: duplicates are detected.", None),
]


def build_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = TRANSACTIONS_SHEET
    ws.append(HEADERS)
    for row in ROWS:
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd"
        row[4].number_format = "#,##0.00"
    for col, width in zip("ABCDEFG", (20, 12, 30, 22, 12, 12, 36)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    info = wb.create_sheet(INSTRUCTIONS_SHEET)
    for left, right in INSTRUCTIONS:
        info.append([left, right])
    info["A1"].font = Font(bold=True, size=14)
    info["A5"].font = Font(bold=True)
    info["B5"].font = Font(bold=True)
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows(min_row=6, max_row=12):
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        row[0].alignment = Alignment(vertical="top")

    return wb


if __name__ == "__main__":
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    build_workbook().save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")
