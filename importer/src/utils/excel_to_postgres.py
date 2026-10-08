import argparse
import json
import logging
import math
from datetime import date, datetime
from decimal import Decimal
from typing import Dict, List

import openpyxl
import xlrd
from dateutil.parser import parse as dateutil_parse

from src.database import db_handler
from src.processing import normalizer
from src.processing.transaction_model import (
    CategorizedTransaction,
    Transaction,
    TransactionSplit,
)

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


def _is_empty(val) -> bool:
    if val is None:
        return True
    if isinstance(val, float) and math.isnan(val):
        return True
    return False


def _to_str(val) -> str:
    return '' if _is_empty(val) else str(val).strip()


def _parse_category(val) -> str | None:
    if _is_empty(val):
        return None
    s = str(val).strip()
    if s.lower() in ('nan', '#n/a', ''):
        return None
    return s


def _to_date(val) -> date | None:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, (int, float)) and val > 0:
        # Excel serial number (openpyxl read_only mode returns dates as numbers)
        from openpyxl.utils.datetime import from_excel
        try:
            return from_excel(val).date()
        except Exception:
            pass
    try:
        return dateutil_parse(str(val)).date()
    except Exception:
        return None


def process_excel_file(excel_file, sheet_name: str = None, column_mapping: Dict[str, str] = None, dry_run: bool = False, file_ext: str = ".xlsx") -> tuple[List[CategorizedTransaction], int, List[dict], List[dict]]:
    """
    Reads an Excel file, transforms the data into CategorizedTransaction objects.
    Uses the first sheet when sheet_name is None.
    Returns (transactions, skipped_count, skip_details, newly_created_accounts).
    newly_created_accounts is a list of {id, name} dicts for accounts created in this run.
    When dry_run=True, no DB writes are made; placeholder IDs are used and newly_created_accounts is [].
    """
    logger.info(f"Processing Excel file: {excel_file}, sheet: {sheet_name or '(first)'}")
    if column_mapping is None:
        column_mapping = {
            'account': 'Account',
            'purchase_date': 'Date',
            'payee': 'Payee',
            'category': 'Category',
            'amount': 'Amount',
            'status': 'Status',
            'description': 'Notes',
        }
    logger.debug(f"Using column mapping: {column_mapping}")

    try:
        if file_ext == ".xls":
            file_data = excel_file.read() if hasattr(excel_file, "read") else open(excel_file, "rb").read()
            wb_xls = xlrd.open_workbook(file_contents=file_data)
            ws_xls = wb_xls.sheet_by_name(sheet_name) if sheet_name else wb_xls.sheet_by_index(0)
            raw_rows = []
            for row_idx in range(ws_xls.nrows):
                row_vals = []
                for col_idx in range(ws_xls.ncols):
                    cell = ws_xls.cell(row_idx, col_idx)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        dt = xlrd.xldate_as_datetime(cell.value, wb_xls.datemode)
                        row_vals.append(dt)
                    elif cell.ctype == xlrd.XL_CELL_EMPTY:
                        row_vals.append(None)
                    else:
                        row_vals.append(cell.value)
                raw_rows.append(tuple(row_vals))
        else:
            wb = openpyxl.load_workbook(excel_file, read_only=True, data_only=True)
            ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
            raw_rows = list(ws.iter_rows(values_only=True))
            wb.close()
    except Exception as e:
        logger.error(f"Error reading Excel file {excel_file} or sheet {sheet_name}: {e}")
        raise

    if not raw_rows:
        return [], 0, [], []

    headers = [_to_str(cell) for cell in raw_rows[0]]
    # _row_num = 1-based Excel row number (row 1 is headers, data starts at 2)
    rows = [dict(zip(headers, row), _row_num=i + 2) for i, row in enumerate(raw_rows[1:])]
    logger.info(f"Successfully read {len(rows)} rows from Excel file.")

    required_fields = ['account', 'purchase_date', 'payee', 'category', 'amount']
    non_null_fields = ['account', 'purchase_date', 'payee', 'amount']

    expected_columns = [column_mapping.get(f) or f for f in required_fields]
    missing_columns = [col for col in expected_columns if col not in headers]
    if missing_columns:
        found_columns = [h for h in headers if h]
        logger.error(f"Missing required columns {missing_columns}; found headers: {found_columns}")
        raise ValueError(
            f"Column headers don't match the expected format. "
            f"Missing: {', '.join(missing_columns)}. "
            f"Expected: {', '.join(expected_columns)}. "
            f"Found in your file: {', '.join(found_columns) or '(none)'}. "
            f"Rename the headers in the first row to match, or download the template in Settings > History."
        )

    initial_rows = len(rows)
    non_null_cols = [column_mapping[f] for f in non_null_fields if f in column_mapping and column_mapping[f] in headers]
    rows_filtered = []
    skipped = 0
    skip_details = []
    for row in rows:
        missing = [col for col in non_null_cols if _is_empty(row.get(col))]
        if not missing:
            rows_filtered.append(row)
        elif all(_to_str(v) == '' for k, v in row.items() if k != '_row_num'):
            continue
        else:
            skip_details.append({"row": row['_row_num'], "reason": f"missing required value: {', '.join(missing)}"})
            skipped += 1

    if skipped > 0:
        logger.warning(f"Skipped {skipped} rows due to missing values in required columns.")
    logger.info(f"Processing {len(rows_filtered)} of {initial_rows} rows after filtering.")

    unique_account_names = list({str(row[column_mapping['account']]).strip() for row in rows_filtered})

    cat_col = column_mapping['category']
    logger.debug(f"Raw category column values before processing: {[row.get(cat_col) for row in rows_filtered]}")
    unique_category_names = list({_parse_category(row.get(cat_col)) for row in rows_filtered} - {None})

    logger.debug(f"Unique account names to process: {unique_account_names}")
    logger.debug(f"Unique category names to process (after filtering NaN/empty): {unique_category_names}")

    if dry_run:
        # Assign placeholder IDs — no DB writes during preview
        account_id_map = {name: 0 for name in unique_account_names}
        newly_created_accounts: List[dict] = []
        category_id_map = {name: 0 for name in unique_category_names}
    else:
        account_id_map, newly_created_accounts = db_handler.get_or_create_accounts_bulk(unique_account_names)
        category_id_map = db_handler.get_or_create_categories_bulk(unique_category_names)
    logger.debug(f"Account ID map: {account_id_map}")
    logger.debug(f"Category ID map: {category_id_map}")

    transactions_to_normalize = []
    original_rows_map = {}

    for row in rows_filtered:
        row_num = row.get('_row_num', '?')
        try:
            account_name = str(row[column_mapping['account']]).strip()
            account_id = account_id_map.get(account_name)
            if account_id is None:
                logger.error(f"Account ID not found for account: {account_name}. Skipping row {row_num}.")
                skip_details.append({"row": row_num, "reason": f"account not found: {account_name}"})
                skipped += 1
                continue

            raw_date = row[column_mapping['purchase_date']]
            purchase_date = _to_date(raw_date)
            if purchase_date is None:
                logger.error(f"Missing valid purchase date for row {row_num}. Skipping.")
                skip_details.append({"row": row_num, "reason": f"invalid date: {raw_date!r}"})
                skipped += 1
                continue

            amount_decimal = Decimal(str(row[column_mapping['amount']]))
            notes = _to_str(row.get(column_mapping['description']))

            transaction_split = TransactionSplit(
                amount=amount_decimal,
                notes=notes
            )

            # 'ignored' is not a review status — it sets the orthogonal is_ignored
            # flag (excluded from reports, still counted in balances) and leaves the
            # review status at its default Pending.
            status_mapping = {
                'pending': 1,
                'cleared': 2,
                'reconciled': 2,
                'ignored': 1
            }
            status_str = _to_str(row.get(column_mapping.get('status', 'pending'))).lower()
            status_id = status_mapping.get(status_str, 1)
            is_ignored = status_str == 'ignored'

            transaction = Transaction(
                purchase_date=purchase_date,
                payment_date=purchase_date,
                payee=str(row[column_mapping['payee']]).strip(),
                account_id=account_id,
                status_id=status_id,
                is_ignored=is_ignored,
                splits=[transaction_split]
            )
            transactions_to_normalize.append(transaction)
            original_rows_map[id(transaction)] = row

        except Exception as e:
            logger.error(f"Error preparing transaction from row {row_num}: {e}")
            skip_details.append({"row": row_num, "reason": f"parse error: {e}"})
            skipped += 1
            continue

    logger.info(f"Normalizing {len(transactions_to_normalize)} transactions...")
    normalized_transactions = normalizer.normalize_transactions(transactions_to_normalize)
    logger.info("Finished normalization.")

    categorized_transactions = []
    for tx in normalized_transactions:
        original_row = original_rows_map.get(id(tx))
        if original_row is None:
            logger.error(f"Could not find original row for normalized transaction: {tx}. Skipping.")
            continue

        category_name = _parse_category(original_row.get(column_mapping['category']))
        category_id = category_id_map.get(category_name) if category_name else None

        categorized_transactions.append(
            CategorizedTransaction(
                transaction=tx,
                category_id=category_id,
                category_name=category_name,
            )
        )
        logger.debug(f"Categorized transaction: {tx.payee} - {tx.amount} (Hash: {tx.unique_hash[:8]}...)")

    if skipped > 0:
        logger.warning(f"Skipped {skipped} rows during transaction preparation.")
    logger.info(f"Finished processing Excel file. Generated {len(categorized_transactions)} categorized transactions.")
    return categorized_transactions, skipped, skip_details, newly_created_accounts


def import_excel_to_db(excel_file: str, sheet_name: str = None, column_mapping: Dict[str, str] = None):
    """
    Main function to process an Excel file and insert the data into the database.
    """
    logger.info(f"Starting import of Excel data to database for file: {excel_file}")
    categorized_transactions, skipped, _, _newly_created = process_excel_file(excel_file, sheet_name, column_mapping)
    logger.info(f"Inserting {len(categorized_transactions)} transactions into the database ({skipped} skipped).")
    db_handler.insert_transactions(categorized_transactions)
    logger.info("Successfully inserted transactions into the database.")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Import Excel data into PostgreSQL database.")
    parser.add_argument("--file", required=True, help="Path to the Excel file.")
    parser.add_argument("--sheet", default="Sheet1", help="Name of the sheet to read from (default: Sheet1).")
    parser.add_argument("--column_mapping", type=json.loads, help='JSON string for column mapping (e.g., {"account": "Conta"}).')

    args = parser.parse_args()

    try:
        import_excel_to_db(args.file, args.sheet, args.column_mapping)
        logger.info(f"Successfully imported data from {args.file} (sheet: {args.sheet}) to the database.")
    except Exception as e:
        logger.error(f"Error importing data: {e}")
